from __future__ import annotations

import csv
import tempfile
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .base import SourceAdapter


class NseDayDeliveryAdapter(SourceAdapter):
    name = "nse_day_delivery"
    timeframe = "day"

    def discover(self) -> Iterable[Path]:
        return sorted(self.root.glob("delivery_*.csv"))

    def ingest(self, connection, source: Path, source_key: str) -> set[str]:
        csv_source = source
        temporary = None
        with source.open("rb") as handle:
            is_excel = handle.read(4) == b"PK\x03\x04"
        if is_excel:
            temporary = tempfile.TemporaryDirectory(prefix="nse-day-")
            csv_source = Path(temporary.name) / "converted.csv"
            with source.open("rb") as excel_input:
                workbook = load_workbook(excel_input, read_only=True, data_only=True)
                try:
                    sheet = workbook.active
                    with csv_source.open("w", newline="", encoding="utf-8") as output:
                        writer = csv.writer(output)
                        writer.writerows(sheet.iter_rows(values_only=True))
                finally:
                    workbook.close()
        # normalize_names handles the spaces in the upstream CSV header.
        try:
            connection.execute(
                """
                INSERT INTO candles
                SELECT
                    ? AS source_key,
                    'day' AS timeframe,
                    trim(symbol) AS symbol,
                    CAST(strptime(trim(date1), '%d-%b-%Y') AS TIMESTAMP) AS timestamp,
                    CAST(open_price AS DOUBLE) AS open,
                    CAST(high_price AS DOUBLE) AS high,
                    CAST(low_price AS DOUBLE) AS low,
                    CAST(close_price AS DOUBLE) AS close,
                    CAST(ttl_trd_qnty AS BIGINT) AS volume,
                    try_cast(deliv_qty AS BIGINT) AS delivery,
                    try_cast(deliv_per AS DOUBLE) AS del_percent
                FROM read_csv(?, header=true, all_varchar=true, normalize_names=true)
                WHERE upper(trim(series)) = 'EQ'
                  AND trim(symbol) <> ''
                """,
                [source_key, str(csv_source)],
            )
        finally:
            if temporary is not None:
                temporary.cleanup()
        return {
            row[0]
            for row in connection.execute(
                "SELECT DISTINCT symbol FROM candles WHERE source_key = ?", [source_key]
            ).fetchall()
        }

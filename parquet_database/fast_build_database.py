from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import sys
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import duckdb
from openpyxl import load_workbook


VERSION = 2


class Progress:
    def __init__(self, log_path: Path) -> None:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        self.log = log_path.open("a", encoding="utf-8", buffering=1)
        self.started = time.monotonic()

    def write(self, percent: float, message: str) -> None:
        elapsed = time.monotonic() - self.started
        line = f"[{datetime.now():%Y-%m-%d %H:%M:%S}] [{max(0, min(100, percent)):6.2f}%] [{elapsed:9.1f}s] {message}"
        print(line, flush=True)
        self.log.write(line + "\n")

    def close(self) -> None:
        self.log.close()


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fast incremental NSE CSV/ZIP to Parquet builder")
    parser.add_argument("--config", type=Path, default=Path(__file__).with_name("config.json"))
    parser.add_argument("--refresh", action="store_true")
    parser.add_argument("--verify-only", action="store_true")
    parser.add_argument("--profile", choices=("local", "ci"), default="local")
    return parser.parse_args()


def configuration(path: Path, profile: str) -> dict:
    path = path.resolve()
    result = json.loads(path.read_text(encoding="utf-8"))
    for key in ("database_dir", "day_source", "day_index_source", "minute_source", "minute_index_source"):
        if key not in result:
            continue
        value = Path(result[key])
        result[key] = (path.parent / value).resolve() if not value.is_absolute() else value.resolve()
    result.setdefault("compression", "zstd")
    selected = result.get("profiles", {}).get(profile, {})
    result.update(selected)
    result.setdefault("memory_limit", "96GB" if profile == "local" else "8GB")
    result.setdefault("threads", max(1, (os.cpu_count() or 4) // 2) if profile == "local" else 4)
    result["profile"] = profile
    return result


def fingerprint(path: Path, adapter_version: int) -> str:
    digest = hashlib.sha256()
    digest.update(f"adapter:{adapter_version}\0".encode())
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def cache_name(kind: str, source: Path) -> str:
    digest = hashlib.sha1(str(source.resolve()).encode()).hexdigest()
    return f"{kind}/{digest}.parquet"


def sql_path(path: Path) -> str:
    return str(path).replace("'", "''")


def copy_day_csv(connection, source: Path, output: Path, compression: str) -> None:
    csv_source = source
    temporary = None
    with source.open("rb") as stream:
        excel = stream.read(4) == b"PK\x03\x04"
    if excel:
        temporary = tempfile.TemporaryDirectory(prefix="nse-day-")
        csv_source = Path(temporary.name) / "converted.csv"
        with source.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                with csv_source.open("w", newline="", encoding="utf-8") as target:
                    csv.writer(target).writerows(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_output = output.with_suffix(".tmp.parquet")
    try:
        connection.execute(
            f"""
            COPY (
                SELECT
                    'day'::VARCHAR timeframe,
                    trim(symbol)::VARCHAR symbol,
                    CAST(strptime(trim(date1), '%d-%b-%Y') AS TIMESTAMP) AS "timestamp",
                    CAST(open_price AS DOUBLE) AS open,
                    CAST(high_price AS DOUBLE) AS high,
                    CAST(low_price AS DOUBLE) AS low,
                    CAST(close_price AS DOUBLE) AS close,
                    CAST(ttl_trd_qnty AS BIGINT) AS volume,
                    try_cast(deliv_qty AS BIGINT) AS delivery,
                    try_cast(deliv_per AS DOUBLE) AS del_percent
                FROM read_csv('{sql_path(csv_source)}', header=true, all_varchar=true, normalize_names=true)
                WHERE upper(trim(series))='EQ' AND trim(symbol)<>''
            ) TO '{sql_path(temp_output)}' (FORMAT PARQUET, COMPRESSION {compression.upper()})
            """
        )
        os.replace(temp_output, output)
    finally:
        if temporary:
            temporary.cleanup()


def copy_day_index_csv(connection, source: Path, output: Path, compression: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_output = output.with_suffix(".tmp.parquet")
    connection.execute(
        f"""
        COPY (
            SELECT
                'day'::VARCHAR timeframe,
                upper(trim(index_name))::VARCHAR symbol,
                CAST(try_strptime(trim(index_date), ['%d-%m-%Y', '%d-%b-%Y', '%d/%m/%Y']) AS TIMESTAMP) AS "timestamp",
                coalesce(try_cast(open_index_value AS DOUBLE), try_cast(closing_index_value AS DOUBLE)) AS open,
                coalesce(try_cast(high_index_value AS DOUBLE), try_cast(closing_index_value AS DOUBLE)) AS high,
                coalesce(try_cast(low_index_value AS DOUBLE), try_cast(closing_index_value AS DOUBLE)) AS low,
                try_cast(closing_index_value AS DOUBLE) AS close,
                try_cast(replace(volume, ',', '') AS BIGINT) AS volume,
                NULL::BIGINT AS delivery,
                NULL::DOUBLE AS del_percent
            FROM read_csv('{sql_path(source)}', header=true, all_varchar=true, normalize_names=true)
            WHERE trim(index_name)<>''
              AND try_strptime(trim(index_date), ['%d-%m-%Y', '%d-%b-%Y', '%d/%m/%Y']) IS NOT NULL
              AND try_cast(closing_index_value AS DOUBLE) IS NOT NULL
        ) TO '{sql_path(temp_output)}' (FORMAT PARQUET, COMPRESSION {compression.upper()})
        """
    )
    os.replace(temp_output, output)


def copy_minute_zip(connection, source: Path, output: Path, universe_file: Path | None,
                    compression: str, ticker_pattern: str = r"(\.NC)?\.NSE$") -> None:
    with zipfile.ZipFile(source) as archive:
        members = [member for member in archive.infolist() if not member.is_dir()]
        if len(members) != 1:
            raise ValueError(f"Expected one CSV member in {source.name}; found {len(members)}")
        with tempfile.TemporaryDirectory(prefix="nse-minute-") as temporary:
            csv_source = Path(archive.extract(members[0], temporary))
            output.parent.mkdir(parents=True, exist_ok=True)
            temp_output = output.with_suffix(".tmp.parquet")
            pattern = ticker_pattern.replace("'", "''")
            symbol_sql = f"regexp_replace(trim(ticker), '{pattern}', '')"
            universe_filter = (
                f"{symbol_sql} IN (SELECT symbol FROM read_parquet('{sql_path(universe_file)}'))"
                if universe_file else "TRUE"
            )
            connection.execute(
                f"""
                COPY (
                    SELECT
                        '1m'::VARCHAR timeframe,
                        {symbol_sql}::VARCHAR symbol,
                        strptime(trim(date)||' '||trim(time), '%d/%m/%Y %H:%M:%S') AS "timestamp",
                        CAST(open AS DOUBLE) AS open,
                        CAST(high AS DOUBLE) AS high,
                        CAST(low AS DOUBLE) AS low,
                        CAST(_close AS DOUBLE) AS close,
                        CAST(volume AS BIGINT) AS volume,
                        NULL::BIGINT AS delivery,
                        NULL::DOUBLE AS del_percent
                    FROM read_csv('{sql_path(csv_source)}', header=true, all_varchar=true, normalize_names=true)
                    WHERE {universe_filter}
                ) TO '{sql_path(temp_output)}' (FORMAT PARQUET, COMPRESSION {compression.upper()})
                """
            )
            os.replace(temp_output, output)


def create_universe(connection, day_cache: list[Path], destination: Path, compression: str) -> str:
    paths = [str(path) for path in day_cache]
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp = destination.with_suffix(".tmp.parquet")
    connection.execute(
        f"COPY (SELECT DISTINCT symbol FROM read_parquet(?)) TO '{sql_path(temp)}' "
        f"(FORMAT PARQUET, COMPRESSION {compression.upper()})",
        [paths],
    )
    digest = connection.execute(
        "SELECT md5(string_agg(symbol, ',' ORDER BY symbol)) FROM read_parquet(?)", [str(temp)]
    ).fetchone()[0]
    os.replace(temp, destination)
    return digest


def safe_symbol(symbol: str) -> str:
    return quote(symbol, safe="-_.&()")


def group(symbol: str) -> str:
    first = symbol[:1].lower()
    return first if first.isascii() and first.isalnum() else "_"


def export_outputs(connection, cache_files: list[Path], database: Path, compression: str,
                   progress: Progress, timeframes: set[str], threads: int) -> None:
    labels = ", ".join(sorted(timeframes))
    bucket_count = 128
    progress.write(82, f"Hash-partitioning {labels} into {bucket_count} parallel export buckets")
    bucket_root = Path(tempfile.mkdtemp(prefix="_buckets_", dir=database))
    connection.execute(f"SET threads={threads}")
    connection.execute(
        f"COPY (SELECT hash(symbol) % {bucket_count} AS bucket_id, * FROM read_parquet(?) "
        f"WHERE timeframe IN ({','.join('?' for _ in timeframes)})) TO '{sql_path(bucket_root)}' "
        f"(FORMAT PARQUET, COMPRESSION {compression.upper()}, PARTITION_BY(bucket_id), "
        f"FILENAME_PATTERN 'data_{{i}}')",
        [[str(path) for path in cache_files], *sorted(timeframes)],
    )
    connection.execute(
        "CREATE OR REPLACE TEMP TABLE export_symbols AS SELECT DISTINCT symbol FROM read_parquet(?) "
        f"WHERE timeframe IN ({','.join('?' for _ in timeframes)})",
        [[str(path) for path in cache_files], *sorted(timeframes)],
    )
    symbols = [row[0] for row in connection.execute("SELECT symbol FROM export_symbols ORDER BY symbol").fetchall()]
    connection.execute("CREATE OR REPLACE TEMP TABLE export_ids AS "
                       "SELECT row_number() OVER(ORDER BY symbol)-1 id, symbol FROM export_symbols")
    ids = dict(connection.execute("SELECT id,symbol FROM export_ids").fetchall())
    try:
        connection.execute(f"SET threads={max(1, min(threads, 16))}")
        placed = 0
        for bucket in range(bucket_count):
            pieces = list((bucket_root / f"bucket_id={bucket}").glob("*.parquet"))
            if not pieces:
                continue
            export_root = Path(tempfile.mkdtemp(prefix=f"_export_{bucket}_", dir=database))
            connection.execute(
                f"COPY (SELECT ids.id partition_id, data.timeframe, data.timestamp, data.open, data.high, "
                f"data.low, data.close, data.volume, data.delivery, data.del_percent "
                f"FROM read_parquet(?) data JOIN export_ids ids USING(symbol) ORDER BY data.symbol) "
                f"TO '{sql_path(export_root)}' (FORMAT PARQUET, COMPRESSION {compression.upper()}, "
                f"PARTITION_BY(partition_id,timeframe), FILENAME_PATTERN 'data_{{i}}')",
                [[str(piece) for piece in pieces]],
            )
            bucket_ids = connection.execute(
                f"SELECT id,symbol FROM export_ids WHERE hash(symbol) % {bucket_count}=?", [bucket]
            ).fetchall()
            for partition_id, symbol in bucket_ids:
                destination = database / group(symbol) / safe_symbol(symbol)
                destination.mkdir(parents=True, exist_ok=True)
                for timeframe in timeframes:
                    output_pieces = list((export_root / f"partition_id={partition_id}" /
                                          f"timeframe={timeframe}").glob("*.parquet"))
                    target = destination / f"{timeframe}.parquet"
                    if len(output_pieces) == 1:
                        os.replace(output_pieces[0], target)
                    elif output_pieces:
                        temporary_target = target.with_suffix(".tmp.parquet")
                        connection.execute(
                            f"COPY (SELECT * FROM read_parquet(?)) TO '{sql_path(temporary_target)}' "
                            f"(FORMAT PARQUET, COMPRESSION {compression.upper()})",
                            [[str(piece) for piece in output_pieces]],
                        )
                        os.replace(temporary_target, target)
                placed += 1
            shutil.rmtree(export_root, ignore_errors=True)
            progress.write(86 + 12 * (bucket + 1) / bucket_count,
                           f"Exported bucket {bucket + 1}/{bucket_count}; placed {placed:,}/{len(ids):,} symbols")
    finally:
        shutil.rmtree(bucket_root, ignore_errors=True)


def verify(connection, database: Path, progress: Progress) -> None:
    files = [path for path in database.glob("*/*/*.parquet") if path.parents[1].name != "_cache"]
    if not files:
        raise RuntimeError("No symbol Parquet files found")
    metadata = connection.execute(
        "SELECT sum(num_rows), count(*) FROM parquet_file_metadata(?)", [[str(path) for path in files]]
    ).fetchone()
    progress.write(100, f"Verified {metadata[1]:,} output files containing {metadata[0]:,} rows")


def main() -> int:
    args = arguments()
    config = configuration(args.config, args.profile)
    database: Path = config["database_dir"]
    for key in ("day_source", "minute_source", "minute_index_source"):
        if key not in config:
            continue
        if not config[key].is_dir():
            raise FileNotFoundError(f"Missing {key}: {config[key]}")
    if args.refresh and database.exists():
        shutil.rmtree(database)
    database.mkdir(parents=True, exist_ok=True)
    for pattern in ("_export_*", "_buckets_*"):
        for orphan in database.glob(pattern):
            if orphan.is_dir():
                shutil.rmtree(orphan)
    combined = database / "_combined.parquet"
    if combined.exists():
        combined.unlink()
    temp_directory = database / "_temp"
    if temp_directory.exists():
        shutil.rmtree(temp_directory)
    temp_directory.mkdir()
    progress = Progress(database / "build.log")
    connection = duckdb.connect()
    connection.execute(f"SET memory_limit='{config['memory_limit']}'")
    connection.execute(f"SET threads={int(config['threads'])}")
    connection.execute(f"SET temp_directory='{sql_path(temp_directory)}'")
    connection.execute("SET preserve_insertion_order=false")
    manifest_path = database / "manifest.json"
    manifest_exists = manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_exists else {"version": VERSION, "sources": {}}
    if manifest.get("version") != VERSION or ((database / "state.duckdb").exists() and not manifest_exists):
        progress.write(0, "Migrating from indexed DuckDB state to direct CSV/Parquet cache")
        for child in database.iterdir():
            if child.name != "build.log":
                if child.is_dir(): shutil.rmtree(child)
                else: child.unlink()
        manifest = {"version": VERSION, "sources": {}}
    legacy_fingerprints = manifest.get("fingerprint_scheme") != "sha256-content-v1"
    manifest["fingerprint_scheme"] = "sha256-content-v1"
    if args.verify_only:
        verify(connection, database, progress)
        return 0

    day_sources = sorted(config["day_source"].glob("delivery_*.csv"))
    day_index_sources = sorted(config.get("day_index_source", Path()).glob("ind_close_all_*.csv")) if config.get("day_index_source") else []
    minute_sources = sorted(config["minute_source"].glob("*.zip"))
    minute_index_sources = sorted(config.get("minute_index_source", Path()).glob("*.zip")) if config.get("minute_index_source") else []
    progress.write(0, f"Profile={config['profile']}; DuckDB memory={config['memory_limit']}; "
                      f"threads={config['threads']}; discovered {len(day_sources):,} equity day CSVs, "
                      f"{len(day_index_sources):,} index day CSVs and "
                      f"{len(minute_sources):,} equity minute ZIPs and "
                      f"{len(minute_index_sources):,} index minute ZIPs")
    changed = False
    dirty_timeframes: set[str] = set()
    live_keys: set[str] = set()
    for index, source in enumerate(day_sources, 1):
        key = "day:" + source.name
        live_keys.add(key)
        cache_rel = cache_name("day", source)
        cache = database / "_cache" / cache_rel
        stamp = fingerprint(source, 2)
        record = manifest["sources"].get(key, {})
        if legacy_fingerprints and record and cache.exists():
            record["fingerprint"] = stamp
            action = "cache fingerprint upgraded"
        elif record.get("fingerprint") != stamp or not cache.exists():
            copy_day_csv(connection, source, cache, config["compression"])
            manifest["sources"][key] = {"fingerprint": stamp, "cache": cache_rel}
            changed = True
            dirty_timeframes.add("day")
            action = "processed"
        else:
            action = "cached"
        progress.write(2 + 18 * index / len(day_sources), f"Day {index:,}/{len(day_sources):,} {action}: {source.name}")

    day_cache = [database / "_cache" / manifest["sources"]["day:" + source.name]["cache"] for source in day_sources]
    universe_file = database / "_cache" / "equity_symbols.parquet"
    old_universe = manifest.get("universe")
    universe_hash = create_universe(connection, day_cache, universe_file, config["compression"])
    universe_changed = universe_hash != old_universe
    manifest["universe"] = universe_hash

    for index, source in enumerate(day_index_sources, 1):
        key = "day_idx:" + source.name
        live_keys.add(key)
        cache_rel = cache_name("day_idx", source)
        cache = database / "_cache" / cache_rel
        stamp = fingerprint(source, 2)
        record = manifest["sources"].get(key, {})
        if record.get("fingerprint") != stamp or not cache.exists():
            copy_day_index_csv(connection, source, cache, config["compression"])
            manifest["sources"][key] = {"fingerprint": stamp, "cache": cache_rel}
            changed = True
            dirty_timeframes.add("day")
            action = "processed"
        else:
            action = "cached"
        progress.write(20 + 5 * index / max(1, len(day_index_sources)), f"Index day {index:,}/{len(day_index_sources):,} {action}: {source.name}")

    minute_inputs = (
        [("minute", source, universe_file, r"(\.NC)?\.NSE$", 3) for source in minute_sources]
        + [("minute_idx", source, None, r"\.NSE_IDX$", 1) for source in minute_index_sources]
    )
    for index, (kind, source, symbol_universe, ticker_pattern, adapter_version) in enumerate(minute_inputs, 1):
        key = kind + ":" + source.name
        live_keys.add(key)
        cache_rel = cache_name(kind, source)
        cache = database / "_cache" / cache_rel
        stamp = fingerprint(source, adapter_version)
        record = manifest["sources"].get(key, {})
        requires_universe_refresh = kind == "minute" and universe_changed
        if legacy_fingerprints and record and cache.exists() and not requires_universe_refresh:
            record["fingerprint"] = stamp
            action = "cache fingerprint upgraded"
        elif requires_universe_refresh or record.get("fingerprint") != stamp or not cache.exists():
            copy_minute_zip(connection, source, cache, symbol_universe, config["compression"], ticker_pattern)
            manifest["sources"][key] = {"fingerprint": stamp, "cache": cache_rel}
            changed = True
            dirty_timeframes.add("1m")
            action = "processed"
        else:
            action = "cached"
        progress.write(20 + 60 * index / len(minute_inputs),
                       f"{kind} {index:,}/{len(minute_inputs):,} {action}: {source.name}")

    for key in list(manifest["sources"]):
        if key not in live_keys:
            cached = database / "_cache" / manifest["sources"][key]["cache"]
            if cached.exists(): cached.unlink()
            del manifest["sources"][key]
            changed = True
            dirty_timeframes.add("day" if key.startswith(("day:", "day_idx:")) else "1m")
    cache_files = [database / "_cache" / row["cache"] for row in manifest["sources"].values()]
    output_exists = any(
        path.parents[1].name != "_cache" for path in database.glob("*/*/*.parquet")
    )
    if changed or not output_exists:
        if not output_exists:
            dirty_timeframes = {"day", "1m"}
        export_outputs(connection, cache_files, database, config["compression"], progress,
                       dirty_timeframes, int(config["threads"]))
    else:
        progress.write(98, "No source changes; symbol outputs are already current")
    verify(connection, database, progress)
    # Commit processing state only after outputs verify. If export is interrupted,
    # the old manifest deliberately causes the changed sources to be retried.
    temporary_manifest = manifest_path.with_suffix(".tmp")
    temporary_manifest.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    os.replace(temporary_manifest, manifest_path)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr, flush=True)
        raise SystemExit(1)
